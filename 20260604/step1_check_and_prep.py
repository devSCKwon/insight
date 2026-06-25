import os
import sys
import datetime
import configparser
import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import PatternFill, Font

def main():
    print("=" * 60)
    print(" 1단계: 실데이터 존재 여부 확인 및 대상 폴더 생성 작업을 시작합니다.")
    print("=" * 60)

    # 1. 설정 파일 읽기
    config_file = 'config.ini'
    if not os.path.exists(config_file):
        print(f"[오류] 설정 파일({config_file})이 존재하지 않습니다.")
        sys.exit(1)

    config = configparser.ConfigParser()
    config.read(config_file, encoding='utf-8')

    try:
        data_list_path = config.get('SETTINGS', 'DATA_LIST_PATH')
        root_dir = config.get('SETTINGS', 'ROOT_DIR')
    except Exception as e:
        print(f"[오류] config.ini 파일 설정을 읽는 중 오류가 발생했습니다: {e}")
        sys.exit(1)

    # 2. 분석 대상 엑셀 파일 확인
    if not os.path.exists(data_list_path):
        print(f"[오류] 분석 대상 엑셀 파일({data_list_path})을 찾을 수 없습니다. 경로를 확인해주세요.")
        sys.exit(1)

    print(f"[*] 설정 정보 로드 완료")
    print(f"   - 대상 엑셀 파일: {data_list_path}")
    print(f"   - 실제 파일 루트: {root_dir}")

    # 3. 엑셀 데이터 로드
    try:
        df = pd.read_excel(data_list_path)
    except Exception as e:
        print(f"[오류] 엑셀 파일을 읽는 중 오류가 발생했습니다: {e}")
        sys.exit(1)

    if 'ALLPATH' not in df.columns:
        print("[오류] 엑셀 파일에 'ALLPATH' 컬럼이 존재하지 않습니다.")
        sys.exit(1)

    # 4. 검사 수행 및 결과 기록을 위한 변수 초기화
    exist_status = []
    full_paths = []
    
    total_rows = len(df)
    exist_count = 0
    missing_count = 0

    print(f"\n[*] 파일 존재 여부 검사 시작 (총 {total_rows}건)...")

    for idx, row in df.iterrows():
        allpath = str(row['ALLPATH']).strip()
        
        # ALLPATH가 비어있는 경우 처리
        if not allpath or pd.isna(row['ALLPATH']):
            exist_status.append('X')
            full_paths.append('')
            missing_count += 1
            continue

        # OS에 맞는 경로 구분자로 통일하고 루트 디렉터리와 병합
        normalized_allpath = os.path.normpath(allpath.replace('/', os.sep))
        full_path = os.path.join(root_dir, normalized_allpath)
        full_paths.append(full_path)

        # 실제 파일 존재 여부 판단
        if os.path.isfile(full_path):
            exist_status.append('O')
            exist_count += 1
        else:
            exist_status.append('X')
            missing_count += 1
            
            # 파일이 존재하지 않는 경우, 해당 파일이 위치할 부모 디렉터리 자동 생성
            dest_dir = os.path.dirname(full_path)
            try:
                if dest_dir and not os.path.exists(dest_dir):
                    os.makedirs(dest_dir, exist_ok=True)
            except Exception as e:
                print(f"[경고] 디렉터리 생성 실패 ({dest_dir}): {e}")

    # 결과 데이터프레임 구축
    df['FULL_PATH'] = full_paths
    df['FILE_EXISTS'] = exist_status
    df['COPY_FROM_PATH'] = ''  # 사용자가 복사용 원본 경로를 작성할 필드 (2단계 대비)

    # 저장 파일명 생성 (시각 포함)
    timestamp = datetime.datetime.now().strftime("%Y%m%d_%H%M%S")
    result_filename = f"Result_{timestamp}.xlsx"
    result_filepath = os.path.join(os.path.dirname(data_list_path) or '.', result_filename)

    # 임시 엑셀 저장
    df.to_excel(result_filepath, index=False)

    # 5. openpyxl을 사용한 스타일링 (파일 존재 여부에 따른 색상 적용)
    try:
        wb = load_workbook(result_filepath)
        ws = wb.active

        # 색상 및 폰트 정의 (Excel 표준 스타일과 유사하게 설정)
        # O (존재): 연한 녹색 채우기 (C6EFCE), 진한 녹색 글자 (006100)
        green_fill = PatternFill(start_color="C6EFCE", end_color="C6EFCE", fill_type="solid")
        green_font = Font(name="Calibri", size=11, color="006100", bold=True)

        # X (누락): 연한 빨간색 채우기 (FFC7CE), 진한 빨간색 글자 (9C0006)
        red_fill = PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid")
        red_font = Font(name="Calibri", size=11, color="9C0006", bold=True)

        # 컬럼 위치 찾기 (1-based index)
        headers = [cell.value for cell in ws[1]]
        file_exists_idx = headers.index('FILE_EXISTS') + 1
        allpath_idx = headers.index('ALLPATH') + 1

        # 데이터 행 순회하며 스타일 적용
        for row in range(2, ws.max_row + 1):
            exists_val = ws.cell(row=row, column=file_exists_idx).value
            
            # FILE_EXISTS 및 ALLPATH 열에 스타일 서식 적용
            if exists_val == 'O':
                ws.cell(row=row, column=file_exists_idx).fill = green_fill
                ws.cell(row=row, column=file_exists_idx).font = green_font
                ws.cell(row=row, column=allpath_idx).fill = green_fill
            else:
                ws.cell(row=row, column=file_exists_idx).fill = red_fill
                ws.cell(row=row, column=file_exists_idx).font = red_font
                ws.cell(row=row, column=allpath_idx).fill = red_fill

        wb.save(result_filepath)
        print(f"[*] 스타일 서식 적용 완료")
    except Exception as e:
        print(f"[경고] 엑셀 셀 색칠(스타일 적용) 중 문제가 발생했습니다: {e}")

    print("=" * 60)
    print(f" 1단계 결과 요약")
    print(f" - 생성된 결과 파일: {result_filepath}")
    print(f" - 전체 대상 파일 수: {total_rows}건")
    print(f" - 존재 확인 파일 수 (O): {exist_count}건")
    print(f" - 누락된 파일 수 (X)  : {missing_count}건 (대상 폴더 생성 완료)")
    print("=" * 60)

if __name__ == '__main__':
    main()

import os
import glob
import sys
import datetime
import configparser
import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import PatternFill, Font, Alignment, Border, Side

def main():
    print("=" * 60)
    print(" 2단계: 누락 데이터 필터링 및 입력 엑셀(Missing_Files) 생성 작업을 시작합니다.")
    print("=" * 60)

    # 1. 설정 파일에서 타겟 폴더 정보 확인
    config_file = 'config.ini'
    data_dir = '.' # 기본 경로
    if os.path.exists(config_file):
        config = configparser.ConfigParser()
        config.read(config_file, encoding='utf-8')
        try:
            data_list_path = config.get('SETTINGS', 'DATA_LIST_PATH')
            data_dir = os.path.dirname(data_list_path) or '.'
        except Exception:
            pass

    # 2. 최신 Result_*.xlsx 파일 검색
    search_pattern = os.path.join(data_dir, "Result_*.xlsx")
    result_files = glob.glob(search_pattern)

    if not result_files:
        print(f"[!] '{data_dir}' 폴더 내에서 Result_*.xlsx 파일을 자동으로 찾을 수 없습니다.")
        # 수동 입력 폴백
        user_input = input("-> Result 엑셀 파일의 경로를 직접 입력해주세요: ").strip()
        if not os.path.exists(user_input):
            print(f"[오류] 파일({user_input})이 존재하지 않습니다. 프로그램을 종료합니다.")
            sys.exit(1)
        target_file = user_input
    else:
        # 가장 최근에 수정된 파일 선택
        target_file = max(result_files, key=os.path.getmtime)
        print(f"[*] 최신 결과 파일 자동 탐색 성공: {target_file}")

    # 3. 데이터 로드 및 필터링
    try:
        df = pd.read_excel(target_file)
    except Exception as e:
        print(f"[오류] 파일을 읽는 중 에러가 발생했습니다: {e}")
        sys.exit(1)

    if 'FILE_EXISTS' not in df.columns or 'COPY_FROM_PATH' not in df.columns:
        print("[오류] 파일 내에 'FILE_EXISTS' 또는 'COPY_FROM_PATH' 열이 없습니다. 1단계가 정상 수행되었는지 확인해주세요.")
        sys.exit(1)

    # 존재하지 않는 파일만 필터링 (FILE_EXISTS == 'X')
    missing_df = df[df['FILE_EXISTS'] == 'X'].copy()

    total_missing = len(missing_df)
    if total_missing == 0:
        print("\n[알림] 누락된 파일이 없습니다! 모든 데이터가 존재하므로 2단계 및 3단계를 수행할 필요가 없습니다.")
        sys.exit(0)

    print(f"[*] 누락된 파일 건수: {total_missing}건")

    # 4. 입력용 엑셀 파일 생성
    timestamp = datetime.datetime.now().strftime("%Y%m%d_%H%M%S")
    missing_filename = f"Missing_Files_{timestamp}.xlsx"
    missing_filepath = os.path.join(data_dir, missing_filename)

    # 누락 데이터 저장
    missing_df.to_excel(missing_filepath, index=False)

    # 5. 사용자 입력 가독성 증대를 위한 스타일 시트 꾸미기 (openpyxl)
    try:
        wb = load_workbook(missing_filepath)
        ws = wb.active

        # 폰트, 색상, 정렬 등 정의
        yellow_fill = PatternFill(start_color="FFF2CC", end_color="FFF2CC", fill_type="solid")  # 옅은 노랑 (입력 유도)
        red_fill = PatternFill(start_color="FCE4D6", end_color="FCE4D6", fill_type="solid")    # 옅은 주홍 (주의)
        
        bold_font = Font(name="Malgun Gothic", size=10, bold=True)
        normal_font = Font(name="Malgun Gothic", size=10)
        
        thin_border = Border(
            left=Side(style='thin', color='D9D9D9'),
            right=Side(style='thin', color='D9D9D9'),
            top=Side(style='thin', color='D9D9D9'),
            bottom=Side(style='thin', color='D9D9D9')
        )

        headers = [cell.value for cell in ws[1]]
        copy_idx = headers.index('COPY_FROM_PATH') + 1
        exists_idx = headers.index('FILE_EXISTS') + 1
        allpath_idx = headers.index('ALLPATH') + 1

        # 셀 테두리 및 서식 전반 처리
        for row in range(2, ws.max_row + 1):
            # 1. 누락 경로 표시 (연한 주홍색)
            ws.cell(row=row, column=exists_idx).fill = red_fill
            ws.cell(row=row, column=allpath_idx).fill = red_fill
            
            # 2. 복사용 원본 입력 필드 표시 (연한 노랑색으로 색칠하여 입력 유도)
            input_cell = ws.cell(row=row, column=copy_idx)
            input_cell.fill = yellow_fill
            
            # 모든 셀에 폰트 및 테두리 적용
            for col in range(1, len(headers) + 1):
                cell = ws.cell(row=row, column=col)
                cell.font = normal_font
                cell.border = thin_border
        
        # 헤더 서식 지정
        header_fill = PatternFill(start_color="D9E1F2", end_color="D9E1F2", fill_type="solid") # 연한 파랑
        for col in range(1, len(headers) + 1):
            cell = ws.cell(row=1, column=col)
            cell.fill = header_fill
            cell.font = bold_font
            cell.alignment = Alignment(horizontal="center", vertical="center")

        wb.save(missing_filepath)
        print(f"[*] 입력 서식 및 하이라이트 적용 완료")
    except Exception as e:
        print(f"[경고] 엑셀 스타일 지정 중 오류가 발생했습니다: {e}")

    print("=" * 60)
    print(f" 2단계 결과 요약")
    print(f" - 생성된 입력 대상 엑셀 파일: {missing_filepath}")
    print(f" - 누락 파일 수: {total_missing}건")
    print(f" [가이드라인]")
    print(f"   1. 생성된 '{missing_filepath}' 파일을 엽니다.")
    print(f"   2. 노란색으로 강조된 'COPY_FROM_PATH' 열에 각 누락 파일의 실제 대체할 복사 대상 파일 경로(전체 경로)를 적어주세요.")
    print(f"      예: C:\\Users\\Name\\Downloads\\photo.png")
    print(f"   3. 입력 완료 후 저장하고 3단계를 실행해주세요.")
    print("=" * 60)

if __name__ == '__main__':
    main()

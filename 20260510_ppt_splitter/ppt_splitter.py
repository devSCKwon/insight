import os
import pandas as pd
from pptx import Presentation
from datetime import datetime
import copy
import re

def create_sample_config(config_path):
    """사용자용 설정 샘플 엑셀 파일 생성"""
    data = {
        '페이지범위': ['110, 111, 112', '120-128'],
        '저장파일명': ['보안정책 위반조치', '개인정보 보호활동']
    }
    df = pd.DataFrame(data)
    df.to_excel(config_path, index=False)
    print(f"샘플 설정 파일 생성됨: {config_path}")

def parse_range(range_str, max_slides):
    """'1, 2, 5-10' 형식의 문자열을 슬라이드 인덱스 리스트(0-based)로 변환"""
    indices = set()
    # 쉼표나 세미콜론으로 구분
    parts = re.split(r'[;,]', str(range_str))
    
    for part in parts:
        part = part.strip()
        if not part:
            continue
            
        if '-' in part or '~' in part:
            # 범위 처리 (예: 120-128)
            start, end = re.split(r'[-~]', part)
            s, e = int(start), int(end)
            # 1-based to 0-based
            indices.update(range(s - 1, e))
        else:
            # 단일 페이지 처리
            indices.add(int(part) - 1)
            
    # 유효한 범위로 제한
    return sorted([i for i in indices if 0 <= i < max_slides])

def split_pptx(source_pptx, config_df, output_dir):
    """설정에 따라 PPTX 분할 저장"""
    if not os.path.exists(output_dir):
        os.makedirs(output_dir)

    for _, row in config_df.iterrows():
        range_str = row['페이지범위']
        target_name = row['저장파일명']
        
        if pd.isna(range_str) or pd.isna(target_name):
            continue

        # 원본 로드 (매번 새로 로드하여 슬라이드 삭제에 의한 영향 방지)
        prs = Presentation(source_pptx)
        total_slides = len(prs.slides)
        
        target_indices = parse_range(range_str, total_slides)
        
        if not target_indices:
            print(f"경고: '{target_name}'에 대한 유효한 페이지 범위가 없습니다 ({range_str})")
            continue

        # 유지할 인덱스를 제외한 나머지 슬라이드 삭제
        # 뒤에서부터 삭제해야 인덱스가 꼬이지 않음
        all_indices = list(range(total_slides))
        to_delete = sorted(list(set(all_indices) - set(target_indices)), reverse=True)
        
        for idx in to_delete:
            rId = prs.slides._sldIdLst[idx].rId
            # Presentation 파트에서 해당 슬라이드와의 관계 제거 (용량 최적화 핵심)
            prs.part.drop_rel(rId)
            # 슬라이드 리스트에서 ID 제거
            del prs.slides._sldIdLst[idx]

        output_path = os.path.join(output_dir, f"{target_name}.pptx")
        prs.save(output_path)
        print(f"저장 완료: {output_path} (페이지: {range_str}, 용량 최적화 적용)")

if __name__ == "__main__":
    BASE_DIR = os.path.dirname(os.path.abspath(__file__))
    SOURCE_DIR = os.path.join(BASE_DIR, "source")
    RESULT_DIR = os.path.join(BASE_DIR, "result")
    CONFIG_FILE = os.path.join(BASE_DIR, "split_config.xlsx")

    # 1. 설정 파일이 없으면 샘플 생성
    if not os.path.exists(CONFIG_FILE):
        create_sample_config(CONFIG_FILE)

    # 2. 원본 폴더에서 PPTX 찾기
    source_files = [f for f in os.listdir(SOURCE_DIR) if f.endswith('.pptx') and not f.startswith('~$')]
    
    if not source_files:
        print(f"알림: '{SOURCE_DIR}' 폴더에 원본 PPTX 파일을 넣어주세요.")
    else:
        # 날짜 시퀀스 폴더 생성
        timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
        target_result_dir = os.path.join(RESULT_DIR, timestamp)
        os.makedirs(target_result_dir, exist_ok=True)
        
        # 첫 번째 PPTX 파일을 대상으로 처리
        source_path = os.path.join(SOURCE_DIR, source_files[0])
        print(f"원본 파일 처리 중: {source_path}")
        print(f"결과 저장 폴더: {target_result_dir}")
        
        # 설정 읽기
        try:
            config_df = pd.read_excel(CONFIG_FILE)
            split_pptx(source_path, config_df, target_result_dir)
            print("\n모든 작업이 완료되었습니다.")
        except Exception as e:
            print(f"오류 발생: {e}")

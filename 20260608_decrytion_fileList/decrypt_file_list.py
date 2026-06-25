import os
import sys
import configparser
from datetime import datetime
import pandas as pd
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.utils import get_column_letter

def get_executable_dir():
    """
    현재 실행 파일(또는 스크립트)이 위치한 디렉터리 경로를 반환합니다.
    PyInstaller로 패키징된 경우와 일반 파이썬 실행 환경을 모두 지원합니다.
    """
    if getattr(sys, 'frozen', False):
        # PyInstaller로 패키징된 실행 파일 환경
        return os.path.dirname(sys.executable)
    else:
        # 일반 파이썬 스크립트 실행 환경
        return os.path.dirname(os.path.abspath(__file__))

def read_config():
    """
    config.ini 파일을 읽어 설정값을 반환합니다.
    """
    exe_dir = get_executable_dir()
    config_path = os.path.join(exe_dir, 'config.ini')

    if not os.path.exists(config_path):
        print(f"[오류] 설정 파일({config_path})을 찾을 수 없습니다.")
        print("프로그램과 같은 위치에 config.ini 파일이 존재해야 합니다.")
        input("종료하려면 엔터키를 누르세요...")
        sys.exit(1)

    config = configparser.ConfigParser()
    try:
        config.read(config_path, encoding='utf-8')
    except UnicodeDecodeError:
        # UTF-8 실패 시 cp949(EUC-KR) 시도
        try:
            config.read(config_path, encoding='cp949')
        except Exception as e:
            print(f"[오류] 설정 파일을 읽는 중 오류가 발생했습니다: {e}")
            input("종료하려면 엔터키를 누르세요...")
            sys.exit(1)

    # 설정값 추출 및 검증
    try:
        encrypted_drive_path = config.get('PATH', 'encrypted_drive_path').strip()
        output_drive_path = config.get('PATH', 'output_drive_path').strip()
        encrypted_extension = config.get('FILE', 'encrypted_extension').strip()
    except (configparser.NoSectionError, configparser.NoOptionError) as e:
        print(f"[오류] 설정 파일의 형식이 잘못되었거나 누락된 항목이 있습니다: {e}")
        input("종료하려면 엔터키를 누르세요...")
        sys.exit(1)

    # 확장자 포맷 정제 (예: '4dhYmt2Ei' -> '.4dhYmt2Ei')
    if not encrypted_extension.startswith('.'):
        encrypted_extension = '.' + encrypted_extension

    return encrypted_drive_path, output_drive_path, encrypted_extension

def scan_encrypted_files(scan_dir, target_ext):
    """
    지정된 디렉터리를 재귀적으로 탐색하여 모든 파일의 리스트를 추출하고,
    마지막 암호화 확장자가 일치하는지 여부에 따라 암호화여부(Y/N) 및 원본 정보를 분기 처리합니다.
    """
    file_list = []
    
    if not os.path.exists(scan_dir):
        print(f"[오류] 파일 스캔 경로가 존재하지 않습니다: {scan_dir}")
        return file_list

    print(f"▶ 스캔 시작 경로: {scan_dir}")
    print(f"▶ 기준 암호화 확장자: {target_ext}")
    print("스캔 중입니다. 잠시만 기다려 주세요...")

    count = 0
    # os.walk를 이용한 재귀 탐색
    for root, dirs, files in os.walk(scan_dir):
        for file in files:
            full_path = os.path.join(root, file)
            # 직속 부모 폴더명 추출
            folder_name = os.path.basename(root)
            if not folder_name:
                # 루트 드라이브일 경우 드라이브명 등을 표시
                folder_name = root

            # 대소문자 구분 없이 암호화 파일 여부 확인
            is_encrypted = file.lower().endswith(target_ext.lower())

            if is_encrypted:
                # 암호화 확장자를 제거하여 원본 파일명 복원
                orig_file_name = file[:-len(target_ext)]
                # 원본 파일명에서 다시 원본 확장자 추출
                _, orig_ext = os.path.splitext(orig_file_name)
                encrypt_yn = 'Y'
            else:
                # 암호화되지 않은 파일의 경우 기존 정보 그대로 유지
                orig_file_name = file
                _, orig_ext = os.path.splitext(file)
                encrypt_yn = 'N'

            file_list.append({
                '폴더명': folder_name,
                '전체경로': full_path,
                '원본파일명': orig_file_name,
                '원본확장자': orig_ext,
                '암호화여부': encrypt_yn
            })
            count += 1
            if count % 100 == 0:
                print(f" - {count}개 파일 스캔 완료...")

    print(f"▶ 스캔 완료! 총 {count}개의 파일이 스캔되었습니다.")
    return file_list

def save_to_excel(file_list, output_dir):
    """
    수집된 파일 정보를 엑셀 파일로 포맷팅하여 저장합니다.
    """
    if not file_list:
        print("[정보] 검색된 파일이 없어 엑셀 파일을 생성하지 않습니다.")
        return None

    # 출력 디렉터리 생성
    os.makedirs(output_dir, exist_ok=True)
    
    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    file_name = f"암호화파일리스트_복원정보_{timestamp}.xlsx"
    output_path = os.path.join(output_dir, file_name)

    print(f"▶ 엑셀 파일 생성 중: {output_path}")

    # 데이터프레임 생성
    df = pd.DataFrame(file_list)

    # ExcelWriter 설정 (openpyxl 엔진 사용)
    with pd.ExcelWriter(output_path, engine='openpyxl') as writer:
        df.to_excel(writer, index=False, sheet_name='파일 목록')
        
        # openpyxl 워크시트 객체 획득하여 스타일 적용
        workbook = writer.book
        worksheet = writer.sheets['파일 목록']
        
        # 스타일 정의
        header_font = Font(name='맑은 고딕', size=11, bold=True, color='FFFFFF')
        header_fill = PatternFill(start_color='4F81BD', end_color='4F81BD', fill_type='solid') # 부드러운 파란색
        data_font = Font(name='맑은 고딕', size=10)
        align_center = Alignment(horizontal='center', vertical='center')
        align_left = Alignment(horizontal='left', vertical='center')
        
        thin_side = Side(border_style="thin", color="D3D3D3")
        thin_border = Border(left=thin_side, right=thin_side, top=thin_side, bottom=thin_side)

        # 헤더 스타일 적용
        for col_num, header in enumerate(df.columns, 1):
            cell = worksheet.cell(row=1, column=col_num)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = align_center
            cell.border = thin_border
            
        # 데이터 스타일 및 정렬 적용
        for row_num in range(2, len(df) + 2):
            for col_num in range(1, len(df.columns) + 1):
                cell = worksheet.cell(row=row_num, column=col_num)
                cell.font = data_font
                cell.border = thin_border
                
                # 전체경로(2번째 열)와 원본파일명(3번째 열)은 왼쪽 정렬,
                # 폴더명(1번째), 원본확장자(4번째), 암호화여부(5번째)는 가운데 정렬
                if col_num in [2, 3]:
                    cell.alignment = align_left
                else:
                    cell.alignment = align_center

        # 행 높이 설정
        worksheet.row_dimensions[1].height = 25 # 헤더 행
        for row in range(2, len(df) + 2):
            worksheet.row_dimensions[row].height = 20

        # 열 너비 자동 조정
        for col in worksheet.columns:
            max_len = 0
            col_letter = get_column_letter(col[0].column)
            for cell in col:
                # 한글 및 유니코드 문자의 실제 폭을 고려한 임시 계산
                val_str = str(cell.value or '')
                cell_len = sum(2 if ord(char) > 128 else 1 for char in val_str)
                if cell_len > max_len:
                    max_len = cell_len
            # 약간의 여백(4) 추가하여 너비 지정, 최대 너비는 80으로 제한
            worksheet.column_dimensions[col_letter].width = min(max(max_len + 4, 12), 80)

    print(f"▶ 엑셀 파일 저장이 완료되었습니다: {file_name}")
    return output_path

def main():
    print("==================================================")
    print("      암호화 파일 복원 리스트 추출 도구 v1.0")
    print("==================================================")

    # 1. 설정값 읽기
    encrypted_drive_path, output_drive_path, encrypted_extension = read_config()

    # 2. 파일 스캔
    file_list = scan_encrypted_files(encrypted_drive_path, encrypted_extension)

    # 3. 엑셀 저장
    if file_list:
        save_to_excel(file_list, output_drive_path)
    else:
        print("[알림] 검색된 파일이 없거나 오류가 있어 엑셀 파일을 저장하지 않았습니다.")

    print("==================================================")
    print("작업이 종료되었습니다.")
    # 단독 exe로 실행 시 창이 바로 닫히지 않도록 키 입력 유도
    input("종료하려면 엔터키를 누르세요...")

if __name__ == "__main__":
    main()

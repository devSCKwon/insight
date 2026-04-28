import pandas as pd
from pptx import Presentation
from copy import deepcopy

def duplicate_slide(pres, index):
    """
    템플릿 슬라이드를 복제하여 새로운 슬라이드를 생성합니다.
    (python-pptx는 직접적인 슬라이드 복제 기능을 제공하지 않아 요소를 복사하는 방식 사용)
    """
    template_slide = pres.slides[index]
    slide_layout = template_slide.slide_layout
    new_slide = pres.slides.add_slide(slide_layout)
    
    for shape in template_slide.shapes:
        el = shape.element
        new_el = deepcopy(el)
        new_slide.shapes._spTree.insert_element_before(new_el, 'p:extLst')
    
    return new_slide

def run_automation(excel_file, template_file, output_file):
    # 1. 엑셀 데이터와 PPT 템플릿 로드
    df = pd.read_excel(excel_file)
    prs = Presentation(template_file)
    
    # 2. 엑셀의 각 행(Row) 데이터를 순회하며 슬라이드 생성
    for _, row in df.iterrows():
        # 첫 번째 슬라이드(index 0)를 복제
        new_slide = duplicate_slide(prs, 0)
        
        # 3. 슬라이드 내의 도형(Shape)들 탐색
        for shape in new_slide.shapes:
            # 제목 등의 일반 텍스트 치환
            if shape.has_text_frame:
                if "{부서명}" in shape.text:
                    shape.text = shape.text.replace("{부서명}", str(row['부서명']))
            
            # 테이블 데이터 입력
            if shape.has_table:
                table = shape.table
                # 엑셀 열 이름에 맞춰서 테이블 셀(Row, Col)에 값 입력
                # cell(1, 0)은 2번째 행, 1번째 열을 의미합니다.
                table.cell(1, 1).text = str(row['카테고리'])
                table.cell(1, 2).text = str(row['평가구분'])
                table.cell(1, 5).text = str(row['점검 문항 내용'])

    # 4. 맨 앞에 남아있는 원본 템플릿 슬라이드 삭제
    xml_slides = prs.slides._sldIdLst
    xml_slides.remove(xml_slides[0])
    
    # 5. 최종 결과 저장
    prs.save(output_file)
    print(f"작업 완료! 생성된 파일: {output_file}")

# 실행부
if __name__ == "__main__":
    # 파일 경로 설정 (본인의 파일명으로 변경하세요)
    EXCEL_PATH = '평가지표.xlsx'
    TEMPLATE_PATH = '진단평가시트_20260428.pptx'
    OUTPUT_PATH = 'Final_Report.pptx'
    
    run_automation(EXCEL_PATH, TEMPLATE_PATH, OUTPUT_PATH)